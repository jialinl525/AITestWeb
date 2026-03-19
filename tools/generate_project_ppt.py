from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt


PRIMARY = RGBColor(15, 76, 129)
PRIMARY_DARK = RGBColor(9, 46, 79)
ACCENT = RGBColor(34, 139, 230)
SUCCESS = RGBColor(46, 160, 67)
WARNING = RGBColor(214, 124, 0)
TEXT = RGBColor(32, 43, 54)
MUTED = RGBColor(90, 104, 120)
LIGHT = RGBColor(245, 248, 252)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(210, 220, 232)


def set_background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_header(slide, title, subtitle=None):
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(8.8), Inches(0.7))
    text_frame = title_box.text_frame
    text_frame.clear()
    p = text_frame.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    accent_bar = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE,
        Inches(0.6),
        Inches(1.0),
        Inches(1.6),
        Inches(0.06),
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = ACCENT
    accent_bar.line.fill.background()

    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.08), Inches(11.5), Inches(0.45))
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.clear()
        p = subtitle_frame.paragraphs[0]
        p.text = subtitle
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(10.5)
        p.font.color.rgb = MUTED


def add_footer(slide, text="AITestWeb | 项目功能与架构介绍"):
    footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(11.6), Inches(0.25))
    tf = footer_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(8.5)
    p.font.color.rgb = MUTED
    p.alignment = PP_ALIGN.RIGHT


def add_bullet_block(slide, left, top, width, height, title, bullets, fill_color=LIGHT):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = LINE

    title_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.12), width - Inches(0.4), Inches(0.35))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    body_box = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.55), width - Inches(0.44), height - Inches(0.7))
    body_frame = body_box.text_frame
    body_frame.clear()
    body_frame.word_wrap = True
    body_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE

    for index, bullet in enumerate(bullets):
        paragraph = body_frame.paragraphs[0] if index == 0 else body_frame.add_paragraph()
        paragraph.text = f"• {bullet}"
        paragraph.level = 0
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(11.5)
        paragraph.font.color.rgb = TEXT
        paragraph.space_after = Pt(6)


def add_table_slide(slide, left, top, width, height, headers, rows, column_widths=None):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table

    if column_widths:
        for index, value in enumerate(column_widths):
            table.columns[index].width = value

    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.name = "Microsoft YaHei"
            paragraph.font.size = Pt(11)
            paragraph.font.bold = True
            paragraph.font.color.rgb = WHITE
            paragraph.alignment = PP_ALIGN.CENTER

    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if row_index % 2 else LIGHT
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(10.5)
                paragraph.font.color.rgb = TEXT
                paragraph.alignment = PP_ALIGN.CENTER

    for row in table.rows:
        for cell in row.cells:
            cell.margin_left = Pt(4)
            cell.margin_right = Pt(4)
            cell.margin_top = Pt(3)
            cell.margin_bottom = Pt(3)


def add_box(slide, left, top, width, height, title, lines, fill_color=WHITE, title_color=PRIMARY_DARK):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.line.color.rgb = LINE
    shape.line.width = Pt(1)

    text_frame = shape.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    title_p = text_frame.paragraphs[0]
    title_p.text = title
    title_p.font.name = "Microsoft YaHei"
    title_p.font.size = Pt(13)
    title_p.font.bold = True
    title_p.font.color.rgb = title_color
    title_p.space_after = Pt(8)

    for line in lines:
        p = text_frame.add_paragraph()
        p.text = line
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT
        p.level = 0
        p.space_after = Pt(4)


def connect(slide, x1, y1, x2, y2, color=ACCENT, width=Pt(1.5)):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = color
    line.line.width = width
    return line


def build_presentation(output_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    # Slide 1: Title
    slide = prs.slides.add_slide(blank)
    set_background(slide, PRIMARY_DARK)
    band = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(7.5))
    band.fill.solid()
    band.fill.fore_color.rgb = PRIMARY_DARK
    band.line.fill.background()
    panel = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(0.85), Inches(12.0), Inches(5.7))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(18, 58, 96)
    panel.line.color.rgb = RGBColor(60, 118, 168)

    title_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(10.8), Inches(1.4))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = "AITestWeb\n项目功能与架构介绍"
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = WHITE

    sub_box = slide.shapes.add_textbox(Inches(1.02), Inches(3.15), Inches(10.5), Inches(1.2))
    sub_frame = sub_box.text_frame
    sub_frame.clear()
    for idx, line in enumerate([
        "面向测试进度、CR 跟踪、KPI 模型对比、人员与任务协同的一体化管理平台",
        f"生成时间：{datetime.now():%Y-%m-%d} | 技术栈：Vue 3 + FastAPI + SQLAlchemy + SQLite"
    ]):
        paragraph = sub_frame.paragraphs[0] if idx == 0 else sub_frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(14 if idx == 0 else 11)
        paragraph.font.color.rgb = RGBColor(220, 235, 247)
        paragraph.space_after = Pt(10)

    highlights = slide.shapes.add_textbox(Inches(1.02), Inches(4.45), Inches(10.8), Inches(1.4))
    hf = highlights.text_frame
    hf.clear()
    for idx, line in enumerate([
        "覆盖业务：测试执行、问题闭环、模型评测、人员协作、权限治理",
        "支持角色隔离、CSV 导入、架构演进兼容、前后端分离部署"
    ]):
        paragraph = hf.paragraphs[0] if idx == 0 else hf.add_paragraph()
        paragraph.text = line
        paragraph.font.name = "Microsoft YaHei"
        paragraph.font.size = Pt(12)
        paragraph.font.color.rgb = RGBColor(210, 228, 240)

    # Slide 2: Overview
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "1. 项目定位与业务价值", "本项目的目标是把测试管理从分散表格与人工同步，收敛到统一的可视化工作台。")
    add_bullet_block(slide, Inches(0.6), Inches(1.55), Inches(5.9), Inches(4.8), "项目定位", [
        "统一管理测试任务、CR 缺陷、模型 KPI、人员工作量和日常任务。",
        "通过前后端分离架构，兼顾快速迭代、页面交互和 API 扩展能力。",
        "支持从‘展示面板’演进到‘可运营配置系统’，如 KPI schema 在线维护、权限组治理。",
    ])
    add_bullet_block(slide, Inches(6.75), Inches(1.55), Inches(5.95), Inches(4.8), "业务价值", [
        "让测试进度与质量风险可视化，便于管理层快速判断阶段状态。",
        "把 CR 跟踪与测试任务关联起来，形成从发现到修复的闭环。",
        "把模型评测数据转为可比较、可追踪、可扩展的 KPI 资产。",
        "把人员与任务、工时和重叠风险关联，帮助资源调度与协作分工。",
    ])
    add_footer(slide)

    # Slide 3: Functional landscape
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "2. 功能全景图", "前端共 9 个主视图，围绕五类核心业务域展开。")
    add_box(slide, Inches(0.7), Inches(1.7), Inches(2.3), Inches(1.45), "测试进度", ["任务列表与详情", "阶段进度 L0/L2/L4", "风险与进度可视化"], LIGHT)
    add_box(slide, Inches(3.25), Inches(1.7), Inches(2.3), Inches(1.45), "CR 跟踪", ["缺陷列表与筛选", "CSV 导入", "Build 快速选择"], LIGHT)
    add_box(slide, Inches(5.8), Inches(1.7), Inches(2.3), Inches(1.45), "KPI / Model Ladder", ["模型排行", "Scatter 对比", "Schema 驱动指标"], LIGHT)
    add_box(slide, Inches(8.35), Inches(1.7), Inches(2.3), Inches(1.45), "人员管理", ["成员画像", "工时分配", "权限与用户管理"], LIGHT)
    add_box(slide, Inches(10.9), Inches(1.7), Inches(1.7), Inches(1.45), "Work Tasks", ["统一任务池", "指派与状态", "详情追踪"], LIGHT)
    add_bullet_block(slide, Inches(0.8), Inches(3.55), Inches(12.0), Inches(2.4), "当前版本重点能力", [
        "测试进度列表已优化为显示当前测试阶段，不再并列展示 L0/L2/L4 三列。",
        "Bug 模块支持重复 CR 导入保护规则、pending confirmation 单击选择镜像、批量删除入口。",
        "KPI 模块支持分类导航、模型详情、在线维护 schema 分类与指标定义。",
        "人员模块支持成员详情画像、工作职责/擅长任务/邮箱维护，以及管理员新增/删除人员。",
    ], fill_color=RGBColor(248, 251, 255))
    add_footer(slide)

    # Slide 4: Role matrix
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "3. 角色与权限体系", "权限体系采用‘用户角色 + 权限组’双通道控制，并在前端路由和后端 API 双侧生效。")
    headers = ["角色/账号", "可见页面", "可编辑范围", "典型场景"]
    rows = [
        ["Administrator\n(username=manager)", "全部页面", "人员/权限组维护、测试与 Bug 编辑、KPI schema 维护", "系统管理员、规则维护者"],
        ["Manager 组成员", "全部页面", "测试进度、Bug、部分运营配置", "测试负责人、项目 owner"],
        ["Internal", "全部页面", "只读，不允许增删改", "内部汇报、旁路查看、跨团队同步"],
        ["Viewer", "仅 FR Progress + Model Ladder", "只读", "普通汇报对象、观察者角色"],
    ]
    add_table_slide(
        slide,
        Inches(0.6),
        Inches(1.75),
        Inches(12.0),
        Inches(4.5),
        headers,
        rows,
        [Inches(2.0), Inches(2.5), Inches(3.5), Inches(4.0)],
    )
    add_footer(slide)

    # Slide 5: Test and work management
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "4. 测试进度与任务协同", "测试主线由 Test Progress 与 Work Tasks 两部分组成，一部分面向测试任务，一部分面向其他工作项。")
    add_bullet_block(slide, Inches(0.6), Inches(1.55), Inches(6.0), Inches(4.9), "Test Progress 能力", [
        "创建和编辑测试任务，维护 FR 编号、测试名称、功能描述、配置方式、测试负责人、开发人员。",
        "记录 L0/L2/L4 三阶段 case 统计，并自动计算总进度与通过情况。",
        "列表页显示风险等级、当前测试阶段、总 case 统计、测试负责人和操作入口。",
        "详情页可联动查看关联 Bug，形成任务到问题的闭环链路。",
    ])
    add_bullet_block(slide, Inches(6.75), Inches(1.55), Inches(5.95), Inches(4.9), "Work Tasks 能力", [
        "管理客户支持、自动化开发等非测试工作项。",
        "维护任务编号、类型、摘要、详细描述、时间范围、状态、进度和 assignee。",
        "与人员模块联动，形成统一任务视角，避免测试任务之外的工作不可见。",
        "和 Bug 模块通过 work_task_id 建立关系，支持非测试任务也能挂载 CR。",
    ])
    add_footer(slide)

    # Slide 6: Bug tracking
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "5. CR 缺陷跟踪能力", "Bug 模块不只是 CRUD 页面，而是围绕导入、筛选、更新效率和数据保护规则设计。")
    add_bullet_block(slide, Inches(0.6), Inches(1.55), Inches(5.8), Inches(4.9), "核心功能", [
        "支持 CR 列表查询、状态筛选、负责人筛选、按测试任务或工作任务关联查看。",
        "支持手动新增、编辑、删除以及统计汇总。",
        "支持从 CSV 导入缺陷，适配已有外部数据源。",
        "支持导出 Top N 数据和按任务查看相关 CR。",
    ])
    add_bullet_block(slide, Inches(6.65), Inches(1.55), Inches(6.05), Inches(4.9), "近期增强点", [
        "重复 CR 导入时，若已有 build 已确认，则不覆盖 build，仅更新时间和 available_images。",
        "对于 pending confirmation 的 build，可在列表中快速单击选择 available_images。",
        "删除场景改为顶部集中入口 + 勾选 + 最终确认，降低误删概率。",
        "支持 test_progress_id 与 work_task_id 双关联，覆盖更完整的问题归属场景。",
    ])
    add_footer(slide)

    # Slide 7: KPI
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "6. KPI / Model Ladder 能力", "KPI 模块是本项目最具分析属性的模块，已经从静态页面演进为 schema 驱动的评测平台。")
    add_bullet_block(slide, Inches(0.6), Inches(1.55), Inches(5.95), Inches(4.9), "展示与分析", [
        "模型性能排行（Ladder）用于同类模型之间的纵向排序。",
        "Scatter 用于两个核心指标之间的横向权衡分析。",
        "模型详情页用于查看某个模型的最新版本记录与指标细节。",
        "准确率类指标统一按 0-100 直接存储与显示，避免前后端缩放误差。",
    ])
    add_bullet_block(slide, Inches(6.75), Inches(1.55), Inches(5.95), Inches(4.9), "运营与扩展", [
        "KPI category/metric 由统一 schema 管理，支持扩展新模型类型与新指标。",
        "后端支持 schema 分类 CRUD，管理员可在线维护模型分类与参数定义。",
        "支持 CSV 模板生成与按分类导入 KPI 数据，适合批量导入评测结果。",
        "前端分类区域已优化为紧凑型导航结构，更适合模型类型增多后的场景。",
    ])
    add_footer(slide)

    # Slide 8: Personnel
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "7. 人员管理与资源视角", "人员模块把‘人’纳入系统主模型，而不是只作为文本字段附着在任务上。")
    add_bullet_block(slide, Inches(0.6), Inches(1.55), Inches(6.0), Inches(4.9), "人员画像", [
        "成员详情页展示 Display Name、Username、Email、Overlap Groups、工作职责、擅长任务。",
        "Administrator 可在详情页编辑成员画像信息。",
        "人员页支持管理员新增/删除用户，并结合权限组进行管理。",
        "当前统计口径已调整为：只统计 manager-group 成员，且排除 Administrator。",
    ])
    add_bullet_block(slide, Inches(6.75), Inches(1.55), Inches(5.95), Inches(4.9), "工时与任务协同", [
        "支持 TaskOwnerAllocation，把测试任务工时分配到具体人员。",
        "若无显式分配，可按 test owner 回退分摊。",
        "可统计每个人的任务数、总工时、完成量、时间区间重叠风险。",
        "与 Work Tasks 联动后，人员视角可以看到测试任务与日常任务的统一列表。",
    ])
    add_footer(slide)

    # Slide 9: Architecture overview
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "8. 系统总体架构", "系统采用典型前后端分离架构，页面、接口、数据存储和权限校验边界清晰。")
    add_box(slide, Inches(0.8), Inches(2.0), Inches(2.2), Inches(2.2), "前端层\nVue 3 + Vite + Element Plus + ECharts", [
        "App.vue 侧边导航",
        "router 控制页面访问",
        "stores/auth 维护登录态",
        "api/* 统一调用后端接口",
    ], fill_color=RGBColor(238, 246, 255))
    add_box(slide, Inches(3.55), Inches(2.0), Inches(2.3), Inches(2.2), "API 层\nFastAPI Routers", [
        "test_progress",
        "bugs",
        "kpi",
        "personnel",
        "work_tasks",
    ], fill_color=RGBColor(245, 250, 255))
    add_box(slide, Inches(6.45), Inches(2.0), Inches(2.25), Inches(2.2), "业务与数据层\nSQLAlchemy Models + Schemas", [
        "Pydantic 数据校验",
        "ORM 实体关系",
        "接口 response model",
    ], fill_color=RGBColor(245, 250, 255))
    add_box(slide, Inches(9.25), Inches(2.0), Inches(2.9), Inches(2.2), "存储与配置层\nSQLite + JSON Schema + Startup Migration", [
        "SQLite 主业务库",
        "KPI schema 持久化",
        "启动时自动补兼容字段",
    ], fill_color=RGBColor(238, 246, 255))
    connect(slide, Inches(3.0), Inches(3.1), Inches(3.55), Inches(3.1))
    connect(slide, Inches(5.85), Inches(3.1), Inches(6.45), Inches(3.1))
    connect(slide, Inches(8.7), Inches(3.1), Inches(9.25), Inches(3.1))
    add_footer(slide)

    # Slide 10: Backend and frontend decomposition
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "9. 前后端模块拆解", "各模块按职责拆分，降低页面复杂度并提升 API 可维护性。")
    add_bullet_block(slide, Inches(0.6), Inches(1.55), Inches(6.0), Inches(4.95), "后端分层要点", [
        "main.py 注册全部 router，并在 startup 阶段完成数据库表创建、兼容字段补齐、默认安全数据初始化。",
        "models.py 定义核心实体：TestProgress、Bug、KPIMetric、User、PermissionGroup、WorkTask、TaskOwnerAllocation。",
        "schemas.py 统一请求/响应契约，降低前后端字段漂移风险。",
        "auth.py 通过 X-User-Name / X-User-Role 识别身份，并结合权限组判断可编辑能力。",
    ])
    add_bullet_block(slide, Inches(6.75), Inches(1.55), Inches(5.95), Inches(4.95), "前端分层要点", [
        "views/ 对应业务页面，router/index.js 定义页面入口与访问守卫。",
        "api/ 目录封装后端访问，统一通过 axios 与 headers 传递当前身份。",
        "stores/auth.js 管理登录态、本地缓存、角色与组信息。",
        "图表页面主要集中在 KPI.vue，通过 ECharts 实现排名和散点可视化。",
    ])
    add_footer(slide)

    # Slide 11: Data model
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "10. 核心数据模型", "系统的数据模型围绕‘任务、问题、模型、人员’四条主线展开。")
    add_box(slide, Inches(0.7), Inches(1.8), Inches(2.2), Inches(1.5), "TestProgress", ["测试任务主表", "阶段统计、进度、负责人", "与 Bug 一对多"], RGBColor(238, 246, 255))
    add_box(slide, Inches(3.25), Inches(1.8), Inches(2.2), Inches(1.5), "Bug", ["CR 记录", "关联测试任务或工作任务", "含 build / available_images"], RGBColor(245, 250, 255))
    add_box(slide, Inches(5.8), Inches(1.8), Inches(2.2), Inches(1.5), "KPIMetric", ["模型评测指标", "按 category/model/version 组织", "支持多指标记录"], RGBColor(238, 246, 255))
    add_box(slide, Inches(8.35), Inches(1.8), Inches(2.2), Inches(1.5), "User / Group", ["用户、角色、权限组", "用于页面可见性与编辑能力", "支持默认 Internal / manager"], RGBColor(245, 250, 255))
    add_box(slide, Inches(10.9), Inches(1.8), Inches(1.7), Inches(1.5), "WorkTask", ["非测试任务", "与人员和 Bug 联动"], RGBColor(238, 246, 255))
    connect(slide, Inches(2.9), Inches(2.55), Inches(3.25), Inches(2.55), SUCCESS)
    connect(slide, Inches(4.0), Inches(3.3), Inches(8.9), Inches(3.3), LINE)
    connect(slide, Inches(10.55), Inches(2.55), Inches(10.9), Inches(2.55), SUCCESS)
    add_bullet_block(slide, Inches(0.8), Inches(3.8), Inches(12.0), Inches(2.1), "模型特点", [
        "TestProgress 与 TaskOwnerAllocation 结合后，可把一个测试任务拆分给多位人员并统计各自工时。",
        "Bug 同时支持 test_progress_id 和 work_task_id，使测试问题与其他任务问题都可纳入同一套缺陷治理流程。",
        "User 与 PermissionGroup 解耦，权限不再只靠单一 role，而是可按组进行能力配置。",
    ], fill_color=RGBColor(248, 251, 255))
    add_footer(slide)

    # Slide 12: Flows
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "11. 关键业务流程", "本项目有两类关键流程：日常操作流和批量数据导入流。")
    add_box(slide, Inches(0.7), Inches(1.75), Inches(5.9), Inches(4.8), "A. 日常操作流", [
        "用户登录 -> 前端 stores/auth 持久化身份 -> axios 为请求注入 X-User-Name / X-User-Role。",
        "页面访问先经过 router.beforeEach，限制 Viewer 只看 FR Progress 和 Model Ladder。",
        "后端 auth.py 再次校验身份与 can_edit_test，确保前后端双重拦截。",
        "管理员/负责人执行新增、编辑、删除操作，最终通过 FastAPI + SQLAlchemy 落库。",
    ], fill_color=RGBColor(238, 246, 255))
    add_box(slide, Inches(6.75), Inches(1.75), Inches(5.85), Inches(4.8), "B. 数据导入与治理流", [
        "Bug 支持 CSV 导入，系统会处理重复 CR、build 覆盖保护和 available_images 更新。",
        "KPI 支持按分类生成模板并导入评测数据，适合模型迭代和多批次版本管理。",
        "KPI schema 已从硬编码走向可维护配置，降低新增模型类型时的开发成本。",
        "startup migration 会自动补齐数据库缺失字段，降低版本升级成本。",
    ], fill_color=RGBColor(245, 250, 255))
    add_footer(slide)

    # Slide 13: Deployment and ops
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "12. 部署方式与运行机制", "项目部署成本低，适合本地试用、团队内网部署和后续数据库升级。")
    add_table_slide(
        slide,
        Inches(0.7),
        Inches(1.8),
        Inches(12.0),
        Inches(3.15),
        ["层级", "当前方案", "说明"],
        [
            ["前端", "Vite + Vue 3", "开发环境 npm run dev，生产环境可构建为静态资源部署"],
            ["后端", "FastAPI + Uvicorn", "REST API 服务，支持 Swagger/ReDoc 文档"],
            ["数据库", "SQLite", "轻量、便于快速落地；README 已说明可迁移到 PostgreSQL"],
            ["初始化", "startup 自动建表/补字段/种子用户", "减少手工迁移步骤，提高版本兼容性"],
        ],
        [Inches(1.5), Inches(2.6), Inches(7.9)],
    )
    add_bullet_block(slide, Inches(0.8), Inches(5.25), Inches(12.0), Inches(1.1), "默认运行入口", [
        "Backend：backend/main.py -> uvicorn main:app；Frontend：frontend/package.json -> vite。",
    ], fill_color=RGBColor(248, 251, 255))
    add_footer(slide)

    # Slide 14: Summary
    slide = prs.slides.add_slide(blank)
    set_background(slide)
    add_header(slide, "13. 总结", "AITestWeb 已不仅是一个演示型仪表板，而是逐步形成了可配置、可治理、可扩展的测试运营平台。")
    add_bullet_block(slide, Inches(0.75), Inches(1.7), Inches(12.0), Inches(4.95), "项目亮点总结", [
        "功能闭环完整：测试任务 -> CR 问题 -> KPI 结果 -> 人员与任务协同。",
        "前后端边界清晰：Vue 页面交互、FastAPI 接口、SQLAlchemy 数据模型分层明确。",
        "权限体系实用：Administrator、Manager 组、Internal、Viewer 覆盖了常见组织角色。",
        "扩展能力较强：KPI schema、自动迁移、CSV 导入、统一 API 结构都支持持续迭代。",
        "适合作为团队测试运营平台基础底座，继续向报表、审批、通知、自动化集成方向演进。",
    ], fill_color=RGBColor(238, 246, 255))
    add_footer(slide, "AITestWeb | Generated by GitHub Copilot (GPT-5.4)")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)


def main():
    project_root = Path(__file__).resolve().parents[1]
    output_path = project_root / "AITestWeb_Project_Overview_CN.pptx"
    build_presentation(output_path)
    print(output_path)


if __name__ == "__main__":
    main()